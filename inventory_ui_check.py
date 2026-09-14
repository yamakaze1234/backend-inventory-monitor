"""Frozen UI acceptance entry. Never starts ERP, listeners or senders."""
import json
from pathlib import Path
import tempfile
import tkinter as tk
from PIL import ImageGrab
from inventory_version import APP_TITLE


def run(output_directory):
    if not output_directory:
        raise ValueError('--ui-check requires an output directory')
    output = Path(output_directory).resolve()
    output.mkdir(parents=True, exist_ok=True)
    from inventory_monitor import InventoryService
    from monitor_app_control import Controller
    from test_inventory_ui_local import build_workspace
    from monitor_setup_wizard import SetupWizard
    with tempfile.TemporaryDirectory(prefix='inventory-frozen-ui-') as temporary:
        service = InventoryService(temporary)
        controller = Controller({'workspace': temporary, 'portable': True})
        from inventory_app import start_configured_message_worker
        from unittest.mock import patch
        with patch.object(controller, 'ensure_current_worker') as launch:
            assert start_configured_message_worker(controller) is False
            launch.assert_not_called()
        from inventory_connection_alert import run_once as check_connection
        from unittest.mock import Mock
        alert_service = InventoryService(Path(temporary) / 'connection-alert-check')
        alert_sender = Mock()
        alert_sender.send.return_value = {'status': 'sent'}
        alert_service.connection_update('', now=100)
        assert check_connection(alert_service, alert_sender, now=100) is False
        assert check_connection(alert_service, alert_sender, now=170) is False
        assert check_connection(alert_service, alert_sender, now=470) is True
        assert check_connection(alert_service, alert_sender, now=200) is False
        alert_sender.send.assert_called_once()
        from notification_robots import RobotStore
        store = RobotStore(temporary)
        first = store.save('运营部机器人（示例）', 'https://oapi.dingtalk.com/robot/send?access_token=offline-fixture-1', ['inventory', 'daily', 'messages'])
        second = store.save('采购部机器人（示例）', 'https://oapi.dingtalk.com/robot/send?access_token=offline-fixture-2', ['daily', 'messages'])
        store.set_enabled(second, False)
        root = tk.Tk()
        root.title(APP_TITLE + ' · 离线界面验收')
        root.geometry('1380x900+20+20')
        root.option_add('*Font', ('Microsoft YaHei UI', 10))
        errors = []
        root.report_callback_exception = lambda kind, value, trace: errors.append(str(value))
        try:
            panel = build_workspace(root, service, live=True, dingtalk_controller=controller)
            assert panel.version_label.cget('text') == APP_TITLE
            for module in ('monitor', 'search', 'config', 'settings', 'daily', 'dingtalk'):
                panel.show_module(module)
                root.update()
                ImageGrab.grab(window=int(root.frame(), 16)).save(output / f'frozen-page-{module}.png')
            row = dict(sku='OFFLINE-5070', goods_id='88102', name='示例 RTX 5070 分库产品', stock=0, able=0, purchase=0)
            service.accept_snapshot('offline-catalog', [row], scope='offline-ui:2')
            from inventory_import import write_template
            from openpyxl import load_workbook
            import_path = Path(temporary) / '产品导入.xlsx'
            write_template(import_path)
            workbook = load_workbook(import_path)
            workbook.active.cell(2, 1, '88102')
            workbook.active.cell(2, 3, 10)
            workbook.save(import_path)
            workbook.close()
            panel.show_module('monitor')
            importer = panel.open_import()
            importer.load_file(import_path)
            root.update()
            assert not service.products()
            assert importer.plan['rows'][0]['product']['goods_id'] == '88102'
            ImageGrab.grab(window=int(importer.frame(), 16)).save(output / 'frozen-product-import.png')
            assert importer.confirm()
            assert service.products()[0]['stock_basis'] == 'company_able'
            assert service.products()[0]['threshold'] == 10
            importer.destroy()
            service.save_product(dict(sku='OFFLINE-CPU', goods_id='88103', name='示例 CPU 265K', threshold=20))
            remover = panel.remove_selected()
            assert not remover.checked
            remover.select_all()
            root.update()
            assert len(remover.checked) == 2
            ImageGrab.grab(window=int(remover.frame(), 16)).save(output / 'frozen-remove-watches.png')
            with patch('inventory_remove_ui.messagebox.askyesno', return_value=True):
                assert remover.confirm()
            assert not service.products()
            remover.destroy()
            request = service.request_warehouses(row['sku'], row['goods_id'])
            details = dict(goods_id='88102', complete=True, depots=[
                dict(id='131', name='稀缺货源', stock=20, able=None, purchase=None),
                dict(id='125', name='样品库', stock=1, able=None, purchase=None)])
            service.accept_warehouses(request['request_id'], details, scope='offline-ui:2')
            product = service.save_product(dict(row, threshold=5, stock_basis='warehouse_stock', warehouse_id='131', warehouse_name='稀缺货源'))
            service.accept_snapshot('offline-details', [row], scope='offline-ui:2', journals={}, warehouse_details={row['sku']: details})
            panel.selected_sku = row['sku']
            panel.show_module('daily')
            root.update()
            ImageGrab.grab(window=int(root.frame(), 16)).save(output / 'frozen-weekly-report.png')
            from inventory_report_preview import open_preview
            fixture_path=output/'preview-fixture.md'
            fixture_path.write_text('# 货源报告 · 离线示例\n\n此处检查报告内容，再选择接收机器人。\n',encoding='utf-8')
            preview=open_preview(root,service,fixture_path)
            root.update()
            assert not preview.reviewed.get()
            assert not preview.selected()
            assert preview.send_button.instate(['disabled'])
            ImageGrab.grab(window=int(preview.frame(),16)).save(output/'frozen-report-preview.png')
            preview.close_preview()
            panel.load_product(product)
            panel.show_module('config')
            root.update()
            assert panel.vars['warehouse_id'].get() == '131'
            assert service.status()['products'][0]['monitor_qty'] == 20
            assert len(panel.warehouse_combo.cget('values')) == 2
            ImageGrab.grab(window=int(root.frame(), 16)).save(output / 'frozen-warehouse-settings.png')
            class SetupFixture:
                portable = True
                def initial_rows(self):
                    return []
                def profiles(self):
                    return []
            wizard = SetupWizard(root, controller, service=SetupFixture())
            root.update()
            ImageGrab.grab(window=int(wizard.win.frame(), 16)).save(output / 'frozen-setup-wizard.png')
            wizard.close()
            panel.show_module('settings')
            root.update()
            assert panel.cache_panel.days.get() == '30天前'
            assert str(panel.cache_panel.clean_button.cget('state')) == 'disabled'
            ImageGrab.grab(window=int(root.frame(),16)).save(output/'frozen-cache-cleanup.png')
            panel.robot_panel.toggle(store.list_robots()[0])
            assert not store.list_robots()[0]['enabled']
            panel.robot_panel.toggle(store.list_robots()[0])
            assert store.list_robots()[0]['enabled']
            panel.robot_panel.edit(store.list_robots()[0])
            root.update()
            editor = next(w for w in panel.robot_panel.winfo_children() if isinstance(w, tk.Toplevel))
            ImageGrab.grab(window=int(editor.frame(), 16)).save(output / 'frozen-robot-editor.png')
            editor.destroy()
            from notification_robot_ui import open_robot_manager
            manager = open_robot_manager(root, temporary)
            root.update()
            ImageGrab.grab(window=int(manager.frame(), 16)).save(output / 'frozen-robot-manager.png')
            manager.destroy()
            root.update()
            ImageGrab.grab(window=int(root.frame(), 16)).save(output / 'frozen-settings.png')
            if errors:
                raise RuntimeError('; '.join(errors))
            (output / 'frozen-ui-check.json').write_text(json.dumps(dict(passed=True, pages=6,
                wizard=True, robot_pause_resume=True, robot_editor=True, warehouse_settings=True, warehouse_qty=20, product_import=True, batch_remove=True, report_preview=True, cache_cleanup=True, inventory_only_no_message_worker=True, erp_outage_one_shot=True, live_services=False), ensure_ascii=False, indent=2), encoding='utf-8')
        finally:
            root.destroy()
