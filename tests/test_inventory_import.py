import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from inventory_monitor import InventoryService
from inventory_import import read_sheet, write_template, prepare_import, confirm_import


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.service = InventoryService(self.root)
        self.catalog = [dict(sku='A', goods_id='101', name='技嘉 Z890 主板'),
                        dict(sku='B', goods_id='102', name='技嘉 Z890 主板'),
                        dict(sku='C', goods_id='103', name='CPU 265K')]
        with self.service.db() as db:
            db.executemany('INSERT INTO catalog VALUES (?,?)', [(p['sku'], json.dumps(p)) for p in self.catalog])

    def row(self, goods_id='', name='', threshold='10', **kwargs):
        return dict(line=2, goods_id=goods_id, name=name, threshold=threshold, error='', **kwargs)

    def test_id_over_name_preview_is_readonly_then_atomic_save(self):
        plan = prepare_import(self.service, [self.row('101', name='旧名称', threshold='0'), self.row(name='CPU 265K')])
        self.assertEqual(self.service.products(), [])
        self.assertEqual(plan['rows'][0]['product']['sku'], 'A')
        self.assertIn('goodsid', plan['rows'][0]['note'])
        self.assertEqual(confirm_import(self.service, plan), 2)
        self.assertEqual({p['stock_basis'] for p in self.service.products()}, {'company_able'})
        self.assertEqual(self.service.products()[0]['threshold'], 0)

    def test_no_fallback_for_wrong_id_and_no_guessing_partial_names(self):
        plan = prepare_import(self.service, [self.row('999', 'CPU 265K'), self.row(name='CPU')])
        self.assertTrue(all(r['problem'] for r in plan['rows']))
        with self.assertRaises(ValueError):
            confirm_import(self.service, plan)
        self.assertEqual(self.service.products(), [])

    def test_ambiguous_name_requires_selection(self):
        row = self.row(name='技嘉 Z890 主板')
        plan = prepare_import(self.service, [row])
        self.assertEqual(len(plan['rows'][0]['candidates']), 2)
        self.assertIsNone(plan['rows'][0]['product'])
        row['selected_sku'] = 'B'
        confirm_import(self.service, prepare_import(self.service, [row]))
        self.assertEqual(self.service.products()[0]['goods_id'], '102')

    def test_duplicate_rows_require_explicit_exclusion(self):
        rows = [self.row('103'), self.row(name='CPU 265K', threshold='20')]
        plan = prepare_import(self.service, rows)
        self.assertTrue(all('重复' in r['problem'] for r in plan['rows']))
        rows[1]['excluded'] = True
        self.assertEqual(confirm_import(self.service, prepare_import(self.service, rows)), 1)

    def test_bad_threshold_and_missing_rows_block_entire_batch(self):
        for threshold in ('', '-1', '1.5', '1e2', 'True', '100000001'):
            with self.subTest(threshold=threshold):
                plan = prepare_import(self.service, [self.row('101'), self.row('103', threshold=threshold)])
                with self.assertRaises(ValueError):
                    confirm_import(self.service, plan)
                self.assertEqual(self.service.products(), [])

    def test_update_forces_company_and_preserves_alert_settings(self):
        old = dict(self.catalog[0], threshold=5, enabled=False, stock_basis='warehouse_stock',
                   warehouse_id='131', warehouse_name='稀缺货源', alerts={'low': False}, keywords='原关键词')
        with self.service.db() as db:
            db.execute('INSERT INTO products VALUES (?,?)', ('A', json.dumps(old)))
            db.execute('INSERT INTO snapshots VALUES (?,?)', ('A', json.dumps({'observed_at': 1})))
        plan = prepare_import(self.service, [self.row('101', threshold='20')])
        self.assertEqual(plan['rows'][0]['action'], '更新')
        confirm_import(self.service, plan)
        product = self.service.products()[0]
        self.assertEqual(product['stock_basis'], 'company_able')
        self.assertEqual(product['warehouse_id'], '')
        self.assertEqual(product['threshold'], 20)
        self.assertFalse(product['enabled'])
        self.assertFalse(product['alerts']['low'])
        self.assertEqual(product['keywords'], '原关键词')
        with self.service.db() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM snapshots').fetchone()[0], 0)

    def test_concurrent_settings_and_catalog_changes_require_new_review(self):
        plan = prepare_import(self.service, [self.row('101')])
        self.service.save_product(dict(self.catalog[2], threshold=7))
        with self.assertRaisesRegex(ValueError, '设置已变化'):
            confirm_import(self.service, plan)
        plan = prepare_import(self.service, [self.row('101')])
        with self.service.db() as db:
            db.execute('UPDATE catalog SET data=? WHERE sku=?', (json.dumps(dict(self.catalog[0], name='新名称')), 'A'))
        with self.assertRaisesRegex(ValueError, '匹配已变化'):
            confirm_import(self.service, plan)

    def test_mid_batch_failure_rolls_back_all_writes(self):
        plan = prepare_import(self.service, [self.row('101'), self.row('103')])
        original = self.service._save_product
        def save(db, product):
            if product['sku'] == 'C':
                raise ValueError('模拟保存失败')
            return original(db, product)
        with patch.object(self.service, '_save_product', side_effect=save):
            with self.assertRaises(ValueError):
                confirm_import(self.service, plan)
        self.assertEqual(self.service.products(), [])
        with self.service.db() as db:
            self.assertEqual(self.service.get(db, 'config_revision', 0), 0)

    def test_csv_aliases_bom_gbk_blank_lines_and_invalid_headers(self):
        path = self.root / 'products.csv'
        for encoding in ('utf-8-sig', 'gb18030'):
            path.write_text('goodsid,名称,预警值\n101,,0\n,,\n,CPU 265K,10\n', encoding=encoding)
            rows = read_sheet(path)
            self.assertEqual([r['line'] for r in rows], [2, 4])
            self.assertEqual(rows[0]['goods_id'], '101')
        path.write_text('goodsid,goods_id,预警值\n101,101,10', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '重复'):
            read_sheet(path)

    def test_xlsx_numeric_ids_formulas_precision_and_template(self):
        from openpyxl import load_workbook
        path = self.root / 'template.xlsx'
        write_template(path)
        book = load_workbook(path)
        sheet = book.active
        sheet.append([])
        sheet.cell(2, 1, 101)
        sheet.cell(2, 3, 0)
        sheet.cell(3, 1, '103')
        sheet.cell(3, 3, '=5+5')
        sheet.cell(4, 1, 1234567890123456)
        sheet.cell(4, 3, 10)
        book.save(path)
        book.close()
        rows = read_sheet(path)
        self.assertEqual(rows[0]['goods_id'], '101')
        self.assertEqual(rows[0]['threshold'], '0')
        self.assertIn('公式', rows[1]['error'])
        self.assertIn('精度', rows[2]['error'])


if __name__ == '__main__':
    unittest.main()
