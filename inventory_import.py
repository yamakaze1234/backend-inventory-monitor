"""Read product sheets and prepare reviewable imports without side effects."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from inventory_rules import validate_product

MAX_ROWS = 5000
HEADERS = {
    'goodsid': 'goods_id', 'goods_id': 'goods_id', 'erp编号': 'goods_id',
    '商品编号': 'goods_id', '名称': 'name', '商品名称': 'name', '产品名称': 'name',
    'name': 'name', '预警值': 'threshold', '预警数量': 'threshold', 'threshold': 'threshold',
}


def cell_text(value, *, identity=False):
    if value is None:
        return ''
    if isinstance(value, bool):
        raise ValueError('不能填写布尔值。')
    if isinstance(value, (int, float)):
        if identity and (value != int(value) or abs(value) >= 10**15):
            raise ValueError('goodsid 数字格式可能丢失精度，请设置为文本并填写完整编号。')
        if value == int(value):
            return str(int(value))
    return str(value).strip()


def read_sheet(path):
    path = Path(path)
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('文件超过 10 MB，请拆分后导入。')
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 100:
                raise ValueError('首张工作表最多 5000 行产品、100 列，请清理多余格式或拆分文件。')
            cells = list(sheet.iter_rows())
            data = [[cell.value for cell in row] for row in cells]
            formulas = {i for i, row in enumerate(cells) if any(c.data_type == 'f' for c in row)}
        finally:
            book.close()
    elif path.suffix.lower() == '.csv':
        raw = path.read_bytes()
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = raw.decode('gb18030')
        import io
        data = list(csv.reader(io.StringIO(content)))
        formulas = {i for i, row in enumerate(data) if any(c.strip().startswith('=') for c in row)}
    else:
        raise ValueError('请选择 .xlsx 或 .csv 文件；旧版 .xls 请先另存为 .xlsx。')
    if not data:
        raise ValueError('表格为空。')
    mapping = {}
    for index, header in enumerate(data[0]):
        key = HEADERS.get(str(header or '').strip().casefold())
        if key:
            if key in mapping:
                raise ValueError(f'表头重复：{header}')
            mapping[key] = index
    if 'threshold' not in mapping or not {'goods_id', 'name'}.intersection(mapping):
        raise ValueError('首行需要“goodsid”或“名称”，以及“预警值”表头。')
    if 0 in formulas:
        raise ValueError('表头不能使用公式。')
    rows = []
    for index, cells in enumerate(data[1:], 1):
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        row = dict(line=index + 1, goods_id='', name='', threshold='', error='')
        try:
            for key, column in mapping.items():
                row[key] = cell_text(cells[column] if column < len(cells) else None, identity=key == 'goods_id')
            if index in formulas:
                row['error'] = '含公式，请粘贴为值后重新导入。'
        except (ValueError, OverflowError) as error:
            row['error'] = str(error)
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise ValueError('每次最多导入 5000 行产品。')
    if not rows:
        raise ValueError('表格没有产品记录。')
    return rows


def write_template(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            csv.writer(stream).writerow(['goodsid', '名称', '预警值'])
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.comments import Comment
    book = Workbook()
    sheet = book.active
    sheet.title = '监控产品'
    sheet.append(['goodsid', '名称', '预警值'])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='285EB9')
    sheet.column_dimensions['A'].width = 24
    sheet.column_dimensions['B'].width = 64
    sheet.column_dimensions['C'].width = 16
    sheet['A1'].comment = Comment('建议填写完整 goodsid，按此编号精确匹配。goodsid 和名称至少填一个；同时填写时以 goodsid 为准。', '库存监控')
    sheet['C1'].comment = Comment('填写 0 至 100000000 的整数。导入统一使用公司大库可销数，确认总览后生效。', '库存监控')
    for index in range(2, 502):
        sheet.cell(index, 1).number_format = '@'
    sheet.freeze_panes = 'A2'
    book.save(path)
    book.close()


def preview_rows(rows, catalog, existing):
    by_sku = {p['sku']: p for p in existing}
    results = []
    for row in rows:
        item = dict(row, candidates=[], product=None, previous=None, action='', problem='', note='')
        results.append(item)
        if row.get('excluded'):
            item['action'] = '已排除'
            continue
        if row.get('error'):
            item['problem'] = row['error']
            continue
        goods_id, name = str(row.get('goods_id', '')).strip(), str(row.get('name', '')).strip()
        if not goods_id and not name:
            item['problem'] = '请填写 goodsid 或名称'
            continue
        if goods_id:
            candidates = [p for p in catalog if str(p.get('goods_id', '')) == goods_id]
            automatic = True
        else:
            candidates = [p for p in catalog if p['name'].strip().casefold() == name.casefold()]
            automatic = bool(candidates)
            if not candidates:
                terms = name.casefold().split()
                candidates = [p for p in catalog if all(t in p['name'].casefold() for t in terms)]
        item['candidates'] = candidates
        chosen = next((p for p in candidates if p['sku'] == row.get('selected_sku')), None)
        if chosen is None and len(candidates) == 1 and automatic:
            chosen = candidates[0]
        if chosen is None:
            item['problem'] = '未匹配，请核对编号或名称并刷新 ERP 目录' if not candidates else f'{len(candidates)} 个候选，请选定商品'
            continue
        if not chosen.get('goods_id'):
            item['problem'] = '目录商品缺少 goodsid，请重新采集 ERP'
            continue
        threshold = str(row.get('threshold', '')).strip()
        if not threshold.isascii() or not threshold.isdecimal() or int(threshold) > 100000000:
            item['problem'] = '预警值须为 0 至 100000000 的整数'
            continue
        previous = by_sku.get(chosen['sku'])
        item['previous'] = previous
        item['product'] = validate_product(dict(previous or {}, sku=chosen['sku'],
            goods_id=chosen['goods_id'], name=chosen['name'], threshold=int(threshold),
            stock_basis='company_able', warehouse_id='', warehouse_name=''))
        item['action'] = '更新' if previous else '新增'
        if goods_id and name and name != chosen['name']:
            item['note'] = '名称不同，按 goodsid 对应商品'
    seen = {}
    for item in results:
        if item['product']:
            seen.setdefault(item['product']['sku'], []).append(item)
    for items in seen.values():
        if len(items) > 1:
            for item in items:
                item['problem'] = '重复商品，请保留一行并排除其他行'
    return results


def prepare_import(service, rows):
    with service.db() as db:
        db.execute('BEGIN')
        catalog = [json.loads(r[0]) for r in db.execute('SELECT data FROM catalog')]
        existing = [json.loads(r[0]) for r in db.execute('SELECT data FROM products')]
        return dict(inputs=[dict(r) for r in rows], rows=preview_rows(rows, catalog, existing),
                    revision=service.get(db, 'config_revision', 0), checked_at=service.get(db, 'checked_at', 0))


def confirm_import(service, plan):
    """Revalidate the reviewed identities and save the whole batch atomically."""
    with service.db() as db:
        db.execute('BEGIN IMMEDIATE')
        if service.get(db, 'config_revision', 0) != plan['revision']:
            raise ValueError('产品设置已变化，请刷新总览后重新确认。')
        catalog = [json.loads(r[0]) for r in db.execute('SELECT data FROM catalog')]
        existing = [json.loads(r[0]) for r in db.execute('SELECT data FROM products')]
        current = preview_rows(plan['inputs'], catalog, existing)
        products = [r['product'] for r in current if r['product']]
        if any(r['problem'] for r in current) or not products:
            raise ValueError('请先处理或排除所有异常行，至少保留一个有效产品。')
        if products != [r['product'] for r in plan['rows'] if r['product']]:
            raise ValueError('ERP 商品匹配已变化，请刷新总览后重新确认。')
        for product in products:
            service._save_product(db, product)
    return len(products)
